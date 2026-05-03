"""
Generates a Word document summarising every file in the project,
what it does, its outputs, and what lives on HDFS.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val

    return table


def main():
    doc = Document()

    # ── Title ─────────────────────────────────────────────────────────────
    title = doc.add_heading("DataScribes Project 2 — Complete File Reference", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "This document describes every file and folder in the project repository, "
        "what each script does, what outputs it produces, and which files live on "
        "HDFS on the NYU DataProc cluster vs. locally on disk."
    )

    # ── 1. Project Overview ───────────────────────────────────────────────
    add_heading(doc, "1. Project Overview", level=1)
    doc.add_paragraph(
        "The goal of this project is to generate high-quality, user-focused descriptions "
        "for 200 public datasets (100 from NYC Open Data, 100 from Data.gov) using the "
        "Claude Sonnet LLM via the Anthropic API. The full pipeline runs on PySpark on "
        "the NYU DataProc cluster. Intermediate data is stored on HDFS; final outputs "
        "are downloaded locally for evaluation and reporting."
    )

    doc.add_paragraph("Pipeline stages (in order):")
    stages = [
        "Ingestion & Sampling — fetch metadata and 5 sample rows per dataset",
        "Profiling — compute column-level statistics (type inference, missing counts)",
        "LLM Description Generation — call Claude Sonnet with a structured prompt",
        "Evaluation — ROUGE, METEOR, BERTScore, NDCG retrieval, manual scoring",
    ]
    for s in stages:
        doc.add_paragraph(s, style="List Number")

    # ── 2. HDFS Files ─────────────────────────────────────────────────────
    add_heading(doc, "2. Files on HDFS (NYU DataProc Cluster)", level=1)
    doc.add_paragraph(
        "These files exist only on the DataProc cluster under "
        "hdfs:///user/km6579_nyu_edu/data/. They are not present locally."
    )

    hdfs_rows = [
        (
            "metadata/combined_metadata_with_samples_v2.parquet",
            "200 rows (100 NYC + 100 Data.gov). Output of dataproc_ingest_and_sample_v2.py. "
            "Contains: dataset_id, source, title, original_description, keywords_json, "
            "column_names_json, column_types_raw_json, download_url, landing_page_url, "
            "last_updated, license, sample_rows_json, sample_error, raw_metadata_json.",
        ),
        (
            "profiles/combined_profiles_spark_v2.parquet",
            "200 rows. Output of dataproc_profile_combined.py. Adds sample_row_count, "
            "sample_column_count, and sample_columns_profile_json (per-column type inference "
            "and example values) to the metadata parquet.",
        ),
        (
            "descriptions/generated_descriptions_sonnet.parquet",
            "200 rows, 10 Spark part files. Output of dataproc_generate_descriptions.py. "
            "Contains: dataset_id, source, title, original_description, "
            "generated_description, generation_model (claude-sonnet-4-5), generation_error. "
            "196 successful, 4 failed (3 rate-limit, 1 prompt-too-long).",
        ),
    ]

    add_table(doc, ["HDFS Path", "Description"], hdfs_rows)

    # ── 3. Local Data Files ───────────────────────────────────────────────
    add_heading(doc, "3. Local Data Files (data/)", level=1)

    local_data_rows = [
        (
            "data/metadata/combined_metadata_with_samples.parquet",
            "20 rows (10 NYC + 10 Data.gov). Early v1 pilot run created by the local "
            "ingestion scripts before DataProc. Now obsolete — superseded by the v2 HDFS file.",
        ),
        (
            "data/descriptions/generated_descriptions_sonnet.parquet/",
            "Folder with 10 Snappy-compressed Spark part files + _SUCCESS marker. "
            "Downloaded from HDFS. Identical to generated_descriptions_all_200.csv. "
            "200 rows, 196 successful LLM descriptions.",
        ),
        (
            "data/evaluation_results.csv",
            "200 rows. Per-dataset automatic metrics: success flag, error type, "
            "original/generated word count, sentence count, novelty score (fraction of "
            "new vocabulary), title coverage flag.",
        ),
        (
            "data/evaluation_summary.csv",
            "3 rows (nyc_open_data, data_gov, OVERALL). Aggregated automatic metrics: "
            "success rate, avg word counts, avg sentences, % following 3-5 sentence rule, "
            "avg novelty score, % title coverage.",
        ),
        (
            "data/text_similarity_results.csv",
            "195 rows (datasets with both original and generated descriptions). "
            "Per-dataset ROUGE-1, ROUGE-2, ROUGE-L, METEOR, BERTScore precision/recall/F1.",
        ),
        (
            "data/retrieval_results.csv",
            "16 rows (8 queries x 2 k values). Simulated TF-IDF search over original vs. "
            "generated descriptions. Columns: query, k, ndcg_original, ndcg_generated, delta. "
            "Generated descriptions improve NDCG@10 by +0.012 on average.",
        ),
        (
            "data/full_evaluation_summary.csv",
            "3 rows (nyc_open_data, data_gov, OVERALL). Average ROUGE-1/2/L, METEOR, "
            "BERTScore F1 per source. Overall BERTScore F1 = 0.860.",
        ),
        (
            "data/manual_evaluation_sample.csv",
            "10 rows (5 NYC + 5 Data.gov). Spreadsheet for human scoring with blank columns: "
            "manual_score_readability, manual_score_accuracy, manual_score_usefulness, "
            "manual_notes.",
        ),
        (
            "data/manual_evaluation_sample.xlsx",
            "Same as manual_evaluation_sample.csv in Excel format for easier hand-scoring.",
        ),
    ]

    add_table(doc, ["File", "Description"], local_data_rows)

    # ── 4. Scripts ────────────────────────────────────────────────────────
    add_heading(doc, "4. Scripts (scripts/)", level=1)
    doc.add_paragraph(
        "Scripts are grouped below by pipeline stage. DataProc scripts run on the cluster "
        "via spark-submit; all others run locally with Python."
    )

    # 4a DataProc scripts
    add_heading(doc, "4a. DataProc Scripts (run on cluster with spark-submit)", level=2)

    dataproc_rows = [
        (
            "dataproc_ingest_and_sample_v2.py",
            "Stage 1 — Ingestion & Sampling",
            "Fetches 100 NYC and 100 Data.gov dataset listings. For NYC uses the Socrata "
            "JSON API (?$limit=5) to get sample rows. For Data.gov streams CSV and stops "
            "after 5 rows. Uses 40 Spark partitions. Writes combined_metadata_with_samples_v2.parquet to HDFS.",
        ),
        (
            "dataproc_ingest_and_sample_combined.py",
            "Stage 1 — v1 (superseded)",
            "Earlier version of the ingestion script. Downloaded full CSVs instead of streaming. "
            "Superseded by v2.",
        ),
        (
            "dataproc_ingest_combined_metadata.py",
            "Stage 1 — metadata only",
            "Fetches metadata without sample rows. Used in early development.",
        ),
        (
            "dataproc_profile_combined.py",
            "Stage 2 — Profiling",
            "Reads combined_metadata_with_samples_v2.parquet from HDFS. Applies three Spark UDFs: "
            "sample_row_count, sample_column_count, build_sample_profile_json (type inference + "
            "example values per column). Writes combined_profiles_spark_v2.parquet to HDFS.",
        ),
        (
            "dataproc_generate_descriptions.py",
            "Stage 3 — LLM Generation",
            "Reads combined_profiles_spark_v2.parquet. Uses mapPartitions with 10 partitions "
            "to call claude-sonnet-4-5 via the Anthropic API (key from ANTHROPIC_API_KEY env var). "
            "Builds a structured prompt per dataset (title, description, keywords, column profiles, "
            "sample row). Writes generated_descriptions_sonnet.parquet to HDFS. "
            "Result: 196/200 successful, 4 failures.",
        ),
    ]

    add_table(doc, ["Script", "Stage", "What it does"], dataproc_rows)

    # 4b Local pipeline scripts
    add_heading(doc, "4b. Local Pipeline Scripts", level=2)

    local_script_rows = [
        ("combine_metadata_with_samples.py", "Merges local metadata and sample row parquet files into one combined file."),
        ("combine_sources.py", "Combines NYC and Data.gov metadata parquets into a single file locally."),
        ("enrich_nyc_with_sample_rows.py", "Adds sample rows to the NYC metadata parquet (local, early version)."),
        ("enrich_data_gov_with_sample_rows.py", "Adds sample rows to the Data.gov metadata parquet (local, early version)."),
        ("json_to_parquet_nyc.py", "Converts raw NYC JSON API responses to Parquet using PySpark."),
        ("json_to_parquet_nyc_pandas.py", "Converts raw NYC JSON API responses to Parquet using pandas (no Spark needed)."),
        ("json_to_parquet_data_gov_pandas.py", "Converts raw Data.gov JSON responses to Parquet using pandas."),
        ("json_with_samples_to_parquet_nyc.py", "Converts enriched NYC JSON (with sample rows) to Parquet."),
        ("json_with_samples_to_parquet_data_gov.py", "Converts enriched Data.gov JSON (with sample rows) to Parquet."),
        ("build_data_gov_llm_input.py", "Builds structured LLM input JSON records from Data.gov profiles."),
        ("generate_data_gov_descriptions_v2.py", "Template-based description generator for Data.gov (v2, no LLM)."),
        ("build_final_project_table.py", "Joins profiles + generated descriptions into final_dataset_description_results.parquet."),
        ("build_report_summary.py", "Groups final table by source; exports summary CSV and XLSX."),
        ("build_manual_eval_sample.py", "Pulls 5 NYC + 5 Data.gov rows into a scoring spreadsheet (expects Parquet input)."),
        ("step1_validate_setup.py", "Checks that all required packages and config files are present before running the pipeline."),
    ]

    add_table(doc, ["Script", "What it does"], local_script_rows)

    # 4c Run scripts
    add_heading(doc, "4c. Run Scripts (convenience wrappers)", level=2)
    doc.add_paragraph(
        "These scripts are thin wrappers that import and call the main() function "
        "of their corresponding module. They exist so the pipeline can be triggered "
        "from the command line without using 'python -m'."
    )

    run_rows = [
        ("run_nyc_ingestion.py", "Runs ingestion/nyc_open_data.py"),
        ("run_data_gov_ingestion.py", "Runs ingestion/data_gov.py"),
        ("run_nyc_profile.py", "Runs profiling/profile_nyc_metadata.py"),
        ("run_data_gov_profile.py", "Runs profiling/profile_data_gov_metadata.py"),
        ("run_build_llm_input.py", "Runs llm/build_llm_input.py"),
        ("run_generate_descriptions.py", "Runs llm/generate_descriptions.py (template v1)"),
        ("run_generate_descriptions_v2.py", "Runs llm/generate_descriptions_v2.py (template v2)"),
        ("run_compare_descriptions.py", "Runs evaluation/compare_descriptions.py"),
        ("run_compare_data_gov_descriptions.py", "Runs evaluation/compare_data_gov_descriptions.py"),
    ]

    add_table(doc, ["Script", "What it calls"], run_rows)

    # ── 5. Ingestion Module ───────────────────────────────────────────────
    add_heading(doc, "5. Ingestion Module (ingestion/)", level=1)

    ingestion_rows = [
        (
            "schema.py",
            "Defines the canonical PySpark StructType schema shared across the pipeline. "
            "Fields: dataset_id, source, title, original_description, keywords (array), "
            "column_names (array), column_types_raw (array), download_url, landing_page_url, "
            "record_count_estimate, last_updated, license, sample_rows_json, raw_metadata_json.",
        ),
        (
            "nyc_open_data.py",
            "fetch_nyc_open_data(limit): calls the Socrata catalog API, fetches per-dataset "
            "detail JSON, normalises to common schema. normalize_dataset() maps raw API fields "
            "to the canonical schema.",
        ),
        (
            "data_gov.py",
            "fetch_data_gov_metadata(limit): calls catalog.data.gov/search, parses DCAT metadata. "
            "normalize_dataset() extracts title, description, keywords, CSV download URL.",
        ),
    ]

    add_table(doc, ["File", "Description"], ingestion_rows)

    # ── 6. LLM Module ─────────────────────────────────────────────────────
    add_heading(doc, "6. LLM Module (llm/)", level=1)
    doc.add_paragraph(
        "These are the local/template-based description generators. The real LLM generation "
        "runs on DataProc via scripts/dataproc_generate_descriptions.py."
    )

    llm_rows = [
        (
            "build_llm_input.py",
            "build_summary(row): assembles a structured dict (llm_input_json) from profiling "
            "output — title, description, keywords, column profiles, sample row count. "
            "Reads nyc_open_data_profiles.parquet, writes nyc_open_data_llm_inputs.parquet.",
        ),
        (
            "generate_descriptions.py",
            "Template baseline v1. build_description(llm_input): simple rule-based description "
            "builder. Model tag: template_baseline_v1. "
            "Reads nyc_open_data_llm_inputs.parquet, writes nyc_open_data_generated_descriptions.parquet.",
        ),
        (
            "generate_descriptions_v2.py",
            "Template baseline v2. Improved with domain detection (choose_domain_phrase) and "
            "column bucketing (bucket_columns). Model tag: template_baseline_v2. "
            "Writes nyc_open_data_generated_descriptions_v2.parquet.",
        ),
    ]

    add_table(doc, ["File", "Description"], llm_rows)

    # ── 7. Profiling Module ───────────────────────────────────────────────
    add_heading(doc, "7. Profiling Module (profiling/)", level=1)

    profiling_rows = [
        (
            "profile_nyc_metadata.py",
            "Local pandas profiling for NYC datasets. Functions: profile_sample_rows, "
            "infer_basic_type, build_profile_record. "
            "Output: data/profiles/nyc_open_data_profiles.parquet.",
        ),
        (
            "profile_data_gov_metadata.py",
            "Local pandas profiling for Data.gov datasets. Same structure as NYC profiler. "
            "Output: data/profiles/data_gov_profiles.parquet.",
        ),
    ]

    add_table(doc, ["File", "Description"], profiling_rows)

    # ── 8. Evaluation Module ──────────────────────────────────────────────
    add_heading(doc, "8. Evaluation Module (evaluation/)", level=1)

    eval_rows = [
        (
            "evaluate_llm_descriptions.py",
            "Automatic metrics: success rate, word count, sentence count, novelty score "
            "(fraction of new vocabulary not in original), title coverage. "
            "Reads generated_descriptions_all_200.csv. "
            "Outputs: data/evaluation_results.csv, data/evaluation_summary.csv.",
        ),
        (
            "evaluate_full.py",
            "Full evaluation covering all three proposal categories: "
            "(1) Text similarity — ROUGE-1/2/L, METEOR, BERTScore F1; "
            "(2) Retrieval — TF-IDF NDCG@5 and NDCG@10 across 8 simulated queries; "
            "(3) Qualitative — 10-dataset manual scoring spreadsheet. "
            "Outputs: text_similarity_results.csv, retrieval_results.csv, "
            "manual_evaluation_sample.csv/.xlsx, full_evaluation_summary.csv.",
        ),
        (
            "compare_descriptions.py",
            "Compares original vs. template v1 vs. template v2 descriptions for NYC datasets "
            "by character length. Output: nyc_open_data_description_comparison.parquet.",
        ),
        (
            "compare_data_gov_descriptions.py",
            "Same comparison script for Data.gov datasets. "
            "Output: data_gov_description_comparison.parquet.",
        ),
    ]

    add_table(doc, ["File", "Description"], eval_rows)

    # ── 9. Key Results ────────────────────────────────────────────────────
    add_heading(doc, "9. Key Results", level=1)

    results_rows = [
        ("Total datasets processed", "200 (100 NYC Open Data + 100 Data.gov)"),
        ("Successful LLM descriptions", "196 / 200 (98%)"),
        ("Failed", "4 — all NYC: 3 rate-limit (429), 1 prompt-too-long (218k tokens)"),
        ("LLM model used", "claude-sonnet-4-5 via Anthropic API"),
        ("Avg generated description length", "102 words"),
        ("Avg sentence count", "4.7 (prompt asked for 3-5)"),
        ("% following 3-5 sentence rule", "73.5% overall (NYC: 88.5%, Data.gov: 59%)"),
        ("Novelty score", "0.67 — 67% of vocabulary is new, not copied from original"),
        ("Title coverage", "99% — nearly all descriptions reference the dataset subject"),
        ("ROUGE-1 (overall)", "0.365"),
        ("ROUGE-2 (overall)", "0.119"),
        ("METEOR (overall)", "0.292"),
        ("BERTScore F1 (overall)", "0.860 — high semantic similarity to original"),
        ("NDCG@10 improvement", "+0.012 — generated descriptions improve retrieval over originals"),
    ]

    add_table(doc, ["Metric", "Value"], results_rows)

    # ── 10. Config & Docs ─────────────────────────────────────────────────
    add_heading(doc, "10. Configuration & Documentation", level=1)

    config_rows = [
        (
            "configs/project_config.yaml",
            "Defines storage paths (data/raw, data/metadata, data/profiles, data/descriptions), "
            "execution engine (spark), and output formats.",
        ),
        (
            "requirements.txt",
            "Python dependencies: pyspark==3.5.1, pandas==2.2.2, requests==2.32.3, "
            "pyyaml==6.0.2, pyarrow==16.1.0.",
        ),
        (
            "docs/dataproc_performance_improvements.md",
            "Documents the 5 optimisations made in the k-dataproc-scalling branch: "
            "stream CSV, Socrata JSON API for NYC, 40 partitions, 30s timeout, "
            "explicit schema fix for createDataFrame crash.",
        ),
        (
            "README.md",
            "V1 scope document: sources (NYC + Data.gov), tabular datasets only, "
            "200 datasets, PySpark on NYU DataProc, Parquet storage, one description per dataset.",
        ),
        (
            ".gitignore",
            "Standard Python gitignore. Excludes __pycache__, .env, *.egg-info, etc.",
        ),
    ]

    add_table(doc, ["File", "Description"], config_rows)

    # ── Save ──────────────────────────────────────────────────────────────
    out_path = Path("data/DataScribes_Project2_File_Reference.docx")
    out_path.parent.mkdir(exist_ok=True)
    doc.save(str(out_path))
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
